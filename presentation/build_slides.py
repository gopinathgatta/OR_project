"""Build the supplementary slides (insert into the existing project deck).

    python presentation/build_slides.py      (run after run_pipeline.py; numbers are read from outputs/)
"""
import json
from pathlib import Path

import pandas as pd
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_LABEL_POSITION
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "outputs"
S = json.loads((OUT / "reports" / "summary.json").read_text())
T = OUT / "tables"

TITLE_FONT, BODY_FONT = "Cambria", "Calibri"
INK, INK2, MUTED = RGBColor(0x11, 0x11, 0x11), RGBColor(0x44, 0x44, 0x44), RGBColor(0x80, 0x80, 0x80)
PANEL = RGBColor(0xDC, 0xE6, 0xF2)        # light-blue panel used in the existing deck
NAVY = RGBColor(0x1F, 0x38, 0x64)
BLUE, ORANGE, AQUA, YELLOW = (RGBColor(0x2A, 0x78, 0xD6), RGBColor(0xEB, 0x68, 0x34),
                              RGBColor(0x1B, 0xAF, 0x7A), RGBColor(0xED, 0xA1, 0x00))
GREY = RGBColor(0x89, 0x87, 0x81)
RED = RGBColor(0xD0, 0x3B, 0x3B)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
BLANK = prs.slide_layouts[6]


def n(v, d=0):
    return f"{v:,.{d}f}"


# --------------------------------------------------------------------------- helpers
def slide(title, notes=None):
    s = prs.slides.add_slide(BLANK)
    tb = s.shapes.add_textbox(Inches(0.6), Inches(0.4), Inches(12.1), Inches(0.9))
    tf = tb.text_frame
    tf.margin_left = tf.margin_right = 0
    p = tf.paragraphs[0]
    p.text = title.upper()
    p.font.name, p.font.size, p.font.color.rgb = TITLE_FONT, Pt(36), INK
    if notes:
        s.notes_slide.notes_text_frame.text = notes
    return s


def text(s, x, y, w, h, items, size=16, color=INK2, bold_first=False, bullets=False, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, font=BODY_FONT, space=6):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.05)
    items = [items] if isinstance(items, str) else items
    for i, it in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(space)
        runs = it if isinstance(it, list) else [(it, {})]
        if bullets:
            runs = [("•  ", {})] + runs
        for txt, o in runs:
            r = p.add_run()
            r.text = txt
            r.font.name = o.get("font", font)
            r.font.size = Pt(o.get("size", size))
            r.font.bold = o.get("bold", bold_first and i == 0)
            r.font.italic = o.get("italic", False)
            r.font.color.rgb = o.get("color", color)
    return tb


def box(s, x, y, w, h, fill=PANEL, line=None, shape=MSO_SHAPE.RECTANGLE, label=None, size=14, color=INK,
        bold=False, align=PP_ALIGN.CENTER):
    sh = s.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line
        sh.line.width = Pt(1.25)
    st_el = sh._element.find(qn("p:style"))
    if st_el is not None:
        sh._element.remove(st_el)
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        sh.adjustments[0] = 0.12
    if label is not None:
        tf = sh.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = Inches(0.08)
        lines = label if isinstance(label, list) else [label]
        for i, ln in enumerate(lines):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = align
            r = p.add_run()
            r.text = ln
            r.font.name, r.font.size, r.font.color.rgb = BODY_FONT, Pt(size if i == 0 else size - 2), color
            r.font.bold = bold and i == 0
    return sh


def arrow(s, x1, y1, x2, y2, color=GREY, width=1.75):
    c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = color
    c.line.width = Pt(width)
    ln = c.line._get_or_add_ln()
    tail = ln.makeelement("{http://schemas.openxmlformats.org/drawingml/2006/main}tailEnd",
                          {"type": "triangle", "w": "med", "len": "med"})
    ln.append(tail)
    return c


def stat(s, x, y, w, value, label, color=NAVY):
    text(s, x, y, w, 0.8, [[(value, {"size": 36, "bold": True, "color": color, "font": TITLE_FONT})]])
    text(s, x, y + 0.8, w, 0.6, label, size=13, color=INK2)


def style_chart(ch, legend=True, size=12):
    ch.font.name, ch.font.size = BODY_FONT, Pt(size)
    ch.font.color.rgb = INK2
    ch.has_legend = legend
    if legend:
        ch.legend.position = XL_LEGEND_POSITION.BOTTOM
        ch.legend.include_in_layout = False
        ch.legend.font.size = Pt(size)
    try:
        va = ch.value_axis
        va.major_gridlines.format.line.color.rgb = RGBColor(0xE1, 0xE0, 0xD9)
        va.format.line.fill.background()
        va.tick_labels.font.color.rgb = MUTED
        ch.category_axis.format.line.color.rgb = RGBColor(0xC3, 0xC2, 0xB7)
        ch.category_axis.tick_labels.font.color.rgb = INK2
    except Exception:
        pass


def table(s, x, y, w, df, col_w=None, size=12, header_fill=NAVY, row_h=0.36, right_cols=None):
    rows, cols = df.shape[0] + 1, df.shape[1]
    gt = s.shapes.add_table(rows, cols, Inches(x), Inches(y), Inches(w), Inches(row_h * rows)).table
    for j, c in enumerate(df.columns):
        if col_w:
            gt.columns[j].width = Inches(col_w[j])
        cell = gt.cell(0, j)
        cell.text = str(c)
        cell.fill.solid(); cell.fill.fore_color.rgb = header_fill
        para = cell.text_frame.paragraphs[0]
        para.font.size, para.font.bold, para.font.color.rgb, para.font.name = Pt(size), True, WHITE, BODY_FONT
    for i, row in enumerate(df.itertuples(index=False), start=1):
        for j, v in enumerate(row):
            cell = gt.cell(i, j)
            cell.text = str(v)
            cell.fill.solid()
            cell.fill.fore_color.rgb = WHITE if i % 2 else RGBColor(0xF3, 0xF6, 0xFA)
            para = cell.text_frame.paragraphs[0]
            para.font.size, para.font.color.rgb, para.font.name = Pt(size), INK, BODY_FONT
            if (right_cols is None and j > 0) or (right_cols is not None and j in right_cols):
                para.alignment = PP_ALIGN.RIGHT
            cell.margin_top = cell.margin_bottom = Inches(0.03)
    return gt


# ========================================================================= slides
b, m1, m2 = S["bottlenecks"], S["model1"], S["model2"]
vr = S["validation_repeat_flow"]["course_specific"]
vg = S["validation_repeat_flow"]["global_rate"]
cap = int(S["parameters"]["section_capacity"])

# 1. Why this dataset? -------------------------------------------------------------
s = slide("Why This Dataset?",
          "Insert immediately after the dataset overview slide. Stress that capacity is NOT in the dataset; "
          "it is a model parameter.")
box(s, 0.6, 1.45, 7.3, 5.55, fill=PANEL)
text(s, 0.85, 1.6, 6.9, 0.5, "Suitability criteria", size=20, color=NAVY, bold_first=True)
crit = [
    [("Real linked records: ", {"bold": True, "color": INK}),
     ("464,739 student-course enrolments, 39,364 students, 4,989 courses, 163 programmes (all match the source)", {})],
    [("Course-level results: ", {"bold": True, "color": INK}),
     ("grades on a 0–10 scale give pass, fail and not-evaluated outcomes, so repetition is observable", {})],
    [("Consecutive years 2021 → 2022: ", {"bold": True, "color": INK}),
     ("transition and retake rates can be estimated", {})],
    [("Programme and curriculum year: ", {"bold": True, "color": INK}),
     ("define stage-based progression stocks", {})],
    [("Anonymised and open: ", {"bold": True, "color": INK}),
     ("hashed IDs and a CC BY 4.0 licence (Zenodo 10.5281/zenodo.17239943), so the work is ethical and reproducible", {})],
]
text(s, 0.85, 2.2, 6.9, 4.7, crit, size=17, bullets=True, space=16)
text(s, 8.3, 1.6, 4.5, 0.5, "Limitations disclosed", size=20, color=RED, bold_first=True)
lims = [
    [("No capacity data: ", {"bold": True, "color": INK}), (f"section capacity is a model parameter (default {cap})", {})],
    [("Only 3 years: ", {"bold": True, "color": INK}), ("2018 is not consecutive and is used for trends only", {})],
    [("Hashed courses: ", {"bold": True, "color": INK}), ("no names and no prerequisite links", {})],
    [("Single university: ", {"bold": True, "color": INK}), ("results may not transfer to other institutions", {})],
    [("Built for dropout research: ", {"bold": True, "color": INK}), ("we use it only for flow and capacity, with no dropout prediction", {})],
]
text(s, 8.3, 2.2, 4.5, 4.7, lims, size=16, bullets=True, space=16)

# 2. Literature review & gap -------------------------------------------------------------
s = slide("Literature Review & Research Gap")
lit = pd.DataFrame([
    ["System Dynamics in higher education", "Galbraith (1998); Sterman (2000)", "Aggregate policy view, no course capacity decision"],
    ["Student-flow / Markov models", "Bartholomew (1982)", "Describe flows, do not optimise capacity"],
    ["Curriculum-flow simulation", "Saltzman & Roeder (2012)", "Simulate only, no section optimisation"],
    ["Timetabling & scheduling IP", "Schaerf (1999); Burke & Petrovic (2002); Navabi-Shirazi et al. (2022)", "Treat demand as fixed, no repeat-demand feedback"],
    ["OR in education (survey)", "Johnes (2015)", "Calls for integrated planning models"],
], columns=["Research domain", "Representative work", "Limitation"])
table(s, 0.6, 1.45, 12.1, lit, col_w=[3.4, 4.6, 4.1], size=14, row_h=0.55, right_cols=[])
box(s, 0.6, 5.0, 12.1, 1.9, fill=PANEL, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
text(s, 0.85, 5.1, 11.6, 1.7, [
    [("Research gap: ", {"bold": True, "color": NAVY, "size": 17}),
     ("progression models show how failure and repetition create demand but do not decide capacity; "
      "capacity IPs decide sections but ignore that unmet capacity returns as next year's repeat demand.", {"size": 16})],
    [("This project: ", {"bold": True, "color": NAVY, "size": 17}),
     ("a data-calibrated SD student-flow model coupled to two Integer Programming models through a backlog feedback loop.", {"size": 16})],
], color=INK, space=8)

# 3. Preprocessing results ---------------------------------------------------------------
s = slide("Data Preprocessing Results")
val = pd.DataFrame([[k.replace("_", " ").capitalize(), n(v["expected"]), n(v["observed"]), "✓"]
                    for k, v in S["validation_counts"].items()],
                   columns=["Check", "Published", "Observed", "Match"])
table(s, 0.6, 1.5, 5.6, val, col_w=[2.2, 1.3, 1.3, 0.8], size=14, row_h=0.45)
steps = [
    "ZIPs extracted in memory; hashed IDs read as text",
    "IDs stripped, lower-cased and validated as hexadecimal hashes",
    "0 duplicate records and 0 rows with missing keys found",
    "9.4% missing grades kept as 'not evaluated' (no credit; they re-enrol like fails)",
    "25 of 178 columns kept (LMS, Wi-Fi and demographics dropped)",
    "Student-course trajectories linked 2018 → 2021 → 2022",
    f"Repeat enrolment flagged: {n(S['repeat_burden_2022']['repeat_enrolments'])} in 2022 "
    f"({S['repeat_burden_2022']['share_of_enrolments']:.1%} of enrolments)",
]
text(s, 6.7, 1.45, 6.0, 0.5, "Cleaning pipeline", size=20, color=NAVY, bold_first=True)
for i, st_ in enumerate(steps):
    y = 2.0 + i * 0.68
    box(s, 6.7, y, 0.45, 0.45, fill=NAVY, shape=MSO_SHAPE.OVAL, label=str(i + 1), size=13, color=WHITE, bold=True)
    text(s, 7.3, y - 0.02, 5.4, 0.6, st_, size=15, color=INK, anchor=MSO_ANCHOR.MIDDLE)

# 4. Student progression ---------------------------------------------------------------
fl = pd.read_csv(T / "progression_flows_2021_2022.csv", index_col=0)
s = slide("Student Progression Analysis")
cd = CategoryChartData()
cd.categories = list(fl.index)
for col in ["Progress %", "Repeat %", "Graduated %", "Withdrawn %"]:
    cd.add_series(col.replace(" %", ""), [round(v, 1) for v in fl[col]])
gf = s.shapes.add_chart(XL_CHART_TYPE.BAR_STACKED_100, Inches(0.6), Inches(1.4), Inches(8.0), Inches(5.6), cd)
ch = gf.chart
style_chart(ch)
ch.plots[0].gap_width = 60
ch.plots[0].overlap = 100
for ser, c in zip(ch.series, [BLUE, ORANGE, AQUA, YELLOW]):
    ser.format.fill.solid(); ser.format.fill.fore_color.rgb = c
for ser, col in zip(ch.series, ["Progress %", "Repeat %", "Graduated %", "Withdrawn %"]):
    for i, v in enumerate(fl[col]):
        lab = ser.points[i].data_label
        if v < 4:                      # remove the label entirely for slivers
            d = lab._get_or_add_dLbl()
            for child in list(d)[1:]:
                d.remove(child)
            d.append(d.makeelement(qn("c:delete"), {"val": "1"}))
            continue
        lab.position = XL_LABEL_POSITION.CENTER
        tf = lab.text_frame
        tf.text = f"{v:.0f}%"
        for r in tf.paragraphs[0].runs:
            r.font.size, r.font.color.rgb = Pt(12), WHITE
ch.category_axis.reverse_order = True
ch.has_title = True
ch.chart_title.text_frame.text = "Share of students in each stage, 2021 → 2022"
ch.chart_title.text_frame.paragraphs[0].font.size = Pt(14)
y1 = fl.loc["Year 1"]
text(s, 9.0, 1.5, 3.8, 5.5, [
    [("Year 1 is the critical stage", {"bold": True, "color": NAVY, "size": 18})],
    f"{y1['Withdrawn %']:.0f}% of first-year students leave, against 4–5% in later years.",
    f"Repetition rises with stage, from {fl.loc['Year 1', 'Repeat %']:.0f}% in Year 1 to "
    f"{fl.loc['Year 4+', 'Repeat %']:.0f}% in Year 4+.",
    f"Final year: {fl.loc['Final year', 'Graduated %']:.0f}% graduate and "
    f"{fl.loc['Final year', 'Repeat %']:.0f}% stay with credits pending.",
    f"{S['repeat_burden_2022']['share_of_students']:.0%} of students retook at least one course in 2022.",
], size=15, color=INK2, space=12)

# 5. SD student-flow model ---------------------------------------------------------------
s = slide("System Dynamics Student-Flow Model",
          "Stocks are students per curriculum stage; flows are calibrated from observed 2021-2022 transitions.")
P = pd.read_csv(T / "progression_transition_matrix.csv", index_col=0)
stages = ["Year 1", "Year 2", "Year 3", "Year 4+", "Final year"]
xs = [1.55, 3.85, 6.15, 8.45, 10.75]
box(s, 0.25, 2.6, 0.95, 0.9, fill=WHITE, line=GREY, shape=MSO_SHAPE.CLOUD, label="Intake", size=12)
for i, (st_, x) in enumerate(zip(stages, xs)):
    box(s, x, 2.45, 1.75, 1.2, fill=NAVY, label=[st_, "stock"], size=16, color=WHITE, bold=True)
    rep = P.loc[st_, st_]
    box(s, x + 0.3, 1.45, 1.15, 0.6, fill=PANEL, shape=MSO_SHAPE.ROUNDED_RECTANGLE,
        label=f"repeat {rep:.0%}", size=12)
    arrow(s, x + 1.2, 2.05, x + 1.2, 2.45, color=ORANGE)
    arrow(s, x + 0.55, 2.45, x + 0.55, 2.05, color=ORANGE)
    wd = P.loc[st_, "Withdrawn"]
    arrow(s, x + 0.875, 3.65, x + 0.875, 4.35, color=YELLOW)
    text(s, x + 0.9, 3.75, 1.0, 0.4, f"{wd:.0%}", size=12, color=INK2)
    if i < 4:
        arrow(s, x + 1.75, 3.05, xs[i + 1], 3.05, color=BLUE, width=2.5)
arrow(s, 1.2, 3.05, xs[0], 3.05, color=BLUE, width=2.5)
box(s, 1.55, 4.35, 10.95, 0.55, fill=RGBColor(0xFD, 0xF1, 0xD6), label="Withdrawn (observed aggregate rate, not predicted per student)", size=13)
arrow(s, 12.5, 3.05, 12.95, 3.05, color=AQUA, width=2.5)
text(s, 12.5, 3.15, 0.8, 0.6, ["grad", f"{P.loc['Final year', 'Graduated']:.0%}"], size=12, color=INK2, space=0)
box(s, 0.6, 5.25, 12.1, 1.75, fill=PANEL, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
text(s, 0.85, 5.35, 7.0, 1.6, [
    [("Stock equation  ", {"bold": True, "color": NAVY}), ("Sⱼ(t+1) = Σᵢ Sᵢ(t)·Pᵢⱼ + Iⱼ(t)", {"font": TITLE_FONT})],
    "P = observed 2021→2022 transition matrix; I = intake (new entrants + re-entrants)",
    "Stocks: in progress (Years 1–4+), repeating, approaching graduation (final year)",
], size=15, color=INK, space=6)
text(s, 8.1, 5.35, 4.5, 1.6, [
    [("Validation (split-half)", {"bold": True, "color": NAVY})],
    f"Mean abs. error {S['validation_student_flow_mape']:.1f}% across stocks and exits; "
    "large stocks within 2.5%",
], size=15, color=INK, space=6)

# 6. Course-flow SD + IP loop ---------------------------------------------------------
s = slide("Course Demand Model & SD–IP Feedback Loop")
nodes = {"data": (0.6, 1.6, "Dataset", "rates 2021 → 2022", GREY),
         "sd": (4.6, 1.6, "System Dynamics", "projects demand D(t)", NAVY),
         "ip": (8.6, 1.6, "Integer Programming", "chooses sections x", AQUA),
         "bl": (4.6, 3.55, "Backlog feedback", "shortfall adds demand", NAVY)}
for k, (x, y, a, bb, c) in nodes.items():
    box(s, x, y, 3.4, 1.2, fill=c, shape=MSO_SHAPE.ROUNDED_RECTANGLE, label=[a, bb], size=18, color=WHITE, bold=True)
arrow(s, 4.0, 2.2, 4.6, 2.2, width=2.25)
arrow(s, 8.0, 2.2, 8.6, 2.2, width=2.25)
arrow(s, 10.3, 2.8, 10.3, 4.15, width=2.25)
arrow(s, 10.3, 4.15, 8.0, 4.15, width=2.25)
arrow(s, 6.3, 3.55, 6.3, 2.8, width=2.25)
text(s, 0.6, 5.1, 6.4, 2.0, [
    [("D", {"font": TITLE_FONT}), ("c", {"font": TITLE_FONT, "size": 12}), ("(t) = F", {"font": TITLE_FONT}),
     ("c", {"font": TITLE_FONT, "size": 12}), ("(1+g)ᵗ + ρ", {"font": TITLE_FONT}), ("c", {"font": TITLE_FONT, "size": 12}),
     (" B", {"font": TITLE_FONT}), ("c", {"font": TITLE_FONT, "size": 12}), ("(t)", {"font": TITLE_FONT})],
    [("B", {"font": TITLE_FONT}), ("c", {"font": TITLE_FONT, "size": 12}),
     ("(t+1) = (1−π", {"font": TITLE_FONT}), ("c", {"font": TITLE_FONT, "size": 12}),
     (")·served + unserved", {"font": TITLE_FONT})],
    [("F first-time, ρ retake rate, π pass rate, B backlog of students not yet passed", {"size": 13, "color": INK2})],
], size=18, color=INK, space=8)
p = S["projection"]
stat(s, 7.4, 5.0, 2.7, n(p["demand"]), f"seats demanded in 2023")
stat(s, 10.2, 5.0, 2.7, n(p["repeat"]), f"of them repeat demand ({p['repeat'] / p['demand']:.1%})", color=ORANGE)
text(s, 7.4, 6.45, 5.5, 0.6, f"Repeat-flow validation: R² = {vr['R2']:.2f} (course rates) vs {vg['R2']:.2f} (single global rate)",
     size=13, color=INK2)

# 7. Bottlenecks ---------------------------------------------------------------
bt = pd.read_csv(T / "bottleneck_courses.csv")
s = slide("Capacity Bottleneck Identification")
stat(s, 0.6, 1.4, 2.9, str(b["bottleneck_courses"]), "bottleneck courses (demand ≥ capacity)", color=RED)
stat(s, 3.6, 1.4, 2.9, str(b["at_risk_courses"]), "at-risk courses (≥ 90% utilised)", color=YELLOW)
stat(s, 6.6, 1.4, 2.9, n(b["projected_unmet_demand"]), "students with unmet demand")
stat(s, 9.6, 1.4, 3.1, str(b["extra_sections_to_clear"]), "extra sections needed to clear all")
top = bt.head(8).assign(Course=lambda d: "C-" + d["course_id"].str[:6])
tt = pd.DataFrame({"Course": top["Course"], "Demand": top["demand"].map(lambda v: n(v)),
                   "Repeat": top["demand_repeat"].map(lambda v: n(v)),
                   "Capacity": top["current_capacity"].map(n),
                   "Utilisation": top["utilisation"].map(lambda v: f"{v:.0%}"),
                   "Unmet": top["unmet_demand"].map(lambda v: n(v))})
table(s, 0.6, 3.05, 7.6, tt, col_w=[1.6, 1.2, 1.2, 1.2, 1.2, 1.2], size=13, row_h=0.42)
box(s, 8.6, 3.05, 4.1, 3.95, fill=PANEL, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
text(s, 8.8, 3.15, 3.7, 3.8, [
    [("How a bottleneck is defined", {"bold": True, "color": NAVY, "size": 16})],
    f"Capacity = {cap} seats per section (model parameter)",
    "Modelled sections = ⌈ 2022 enrolment ÷ capacity ⌉",
    "Bottleneck if projected demand ≥ modelled capacity",
    f"Repeat share in bottlenecks: {b['bottleneck_repeat_share']:.0%} vs {p['repeat'] / p['demand']:.0%} overall; repetition drives congestion",
    f"{b['programmes_affected']} programmes affected",
], size=14, color=INK, space=8)

# 8. Model 1 ---------------------------------------------------------------
sc = pd.read_csv(T / "sensitivity_capacity.csv")
s = slide("Optimization 1: Minimum Section Allocation")
box(s, 0.6, 1.4, 5.6, 2.3, fill=PANEL, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
text(s, 0.85, 1.5, 5.2, 2.2, [
    [("Integer Programming", {"bold": True, "color": NAVY, "size": 16})],
    [("min  Σ x", {"font": TITLE_FONT, "size": 26}), ("c", {"font": TITLE_FONT, "size": 15})],
    [("s.t.  cap · x", {"font": TITLE_FONT, "size": 23}), ("c", {"font": TITLE_FONT, "size": 15}),
     ("  ≥  D", {"font": TITLE_FONT, "size": 23}), ("c", {"font": TITLE_FONT, "size": 15}),
     ("     x", {"font": TITLE_FONT, "size": 23}), ("c", {"font": TITLE_FONT, "size": 15}),
     (" ∈ ℤ⁺,  x", {"font": TITLE_FONT, "size": 23}), ("c", {"font": TITLE_FONT, "size": 15}), (" ≥ 1", {"font": TITLE_FONT, "size": 23})],
], color=INK, space=6)
stat(s, 0.6, 3.95, 2.7, n(m1["total_sections"]), "minimum sections meeting all demand")
stat(s, 3.4, 3.95, 2.8, f"+{m1['sections_added']} / −{m1['sections_released']}",
     "sections added / released vs modelled current")
text(s, 0.6, 5.55, 5.6, 1.4, [
    f"LP relaxation: {n(m1['lp_relaxation_sections'], 1)} fractional sections. The {n(m1['integrality_gap_sections'])}-section "
    "integrality gap shows why fractional sections cannot be used.",
    "Solved with HiGHS (optimal); every value is an integer and all demand is met.",
], size=14, color=INK2, space=6)
cd = CategoryChartData()
cd.categories = [str(int(c)) for c in sc["capacity"]]
cd.add_series("Integer solution (Model 1)", list(sc["model1_sections"]))
cd.add_series("LP relaxation", [round(v) for v in sc["lp_relaxation_sections"]])
gf = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(6.6), Inches(1.4), Inches(6.1), Inches(5.6), cd)
ch = gf.chart
style_chart(ch)
for ser, c in zip(ch.series, [BLUE, PANEL]):
    ser.format.fill.solid(); ser.format.fill.fore_color.rgb = c
ch.series[1].format.line.color.rgb = BLUE
ch.plots[0].gap_width = 80
ch.has_title = True
ch.chart_title.text_frame.text = "Total sections vs section-capacity assumption"
ch.chart_title.text_frame.paragraphs[0].font.size = Pt(14)

# 9. Model 2 ---------------------------------------------------------------
sb = pd.read_csv(T / "sensitivity_budget.csv")
s = slide("Optimization 2: Limited Additional Capacity")
box(s, 0.6, 1.4, 5.6, 2.55, fill=PANEL, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
text(s, 0.85, 1.5, 5.2, 2.4, [
    [("Binary Integer Programming", {"bold": True, "color": NAVY, "size": 16})],
    [("max  Σ w·a·y", {"font": TITLE_FONT, "size": 23}),
     ("   y ∈ {0,1}: k-th extra section to course c", {"size": 13, "color": INK2})],
    [("s.t.  Σ y ≤ B   (extra sections available)", {"font": TITLE_FONT, "size": 18})],
    [("y", {"font": TITLE_FONT, "size": 18}), ("c,k", {"font": TITLE_FONT, "size": 13}), (" ≤ y", {"font": TITLE_FONT, "size": 18}),
     ("c,k−1", {"font": TITLE_FONT, "size": 13}), ("  (open in order);  Σ", {"font": TITLE_FONT, "size": 18}),
     ("p", {"font": TITLE_FONT, "size": 13}), (" y ≤ M  (fairness)", {"font": TITLE_FONT, "size": 18})],
], color=INK, space=8)
stat(s, 0.6, 4.15, 2.7, f"{m2['unmet_reduction_pct']:.0f}%", f"unmet demand removed with B = {m2['extra_sections_available']}", color=AQUA)
stat(s, 3.4, 4.15, 2.8, f"{n(m2['unmet_demand_before'])} → {n(m2['unmet_demand_after'])}", "students unmet, before → after")
text(s, 0.6, 5.75, 5.6, 1.3, [
    f"Courses still short: {m2['courses_short_before']} → {m2['courses_short_after']}. "
    f"With a fairness cap of 5 per programme: {n(S['model2_fair']['unmet_demand_after'])} unmet.",
    "The highest-shortage courses are served first, even when the budget cannot cover every course.",
], size=14, color=INK2, space=6)
cd = CategoryChartData()
cd.categories = [str(int(v)) for v in sb["extra_sections"]]
cd.add_series("Unmet demand after allocation", [round(v, 1) for v in sb["unmet_demand_after"]])
gf = s.shapes.add_chart(XL_CHART_TYPE.LINE_MARKERS, Inches(6.6), Inches(1.4), Inches(6.1), Inches(5.6), cd)
ch = gf.chart
style_chart(ch, legend=False)
ser = ch.series[0]
ser.format.line.color.rgb = BLUE; ser.format.line.width = Pt(2.5); ser.smooth = False
ser.marker.format.fill.solid(); ser.marker.format.fill.fore_color.rgb = BLUE
ser.marker.format.line.color.rgb = BLUE
ch.has_title = True
ch.chart_title.text_frame.text = "Unmet demand (students) vs additional sections B"
ch.chart_title.text_frame.paragraphs[0].font.size = Pt(14)

# 10. Scenario comparison ------------------------------------------------------------
comp = pd.read_csv(T / "scenario_comparison.csv")
s = slide("Scenario Comparison (2023–2027)")
cd = CategoryChartData()
cd.categories = list(comp["scenario"])
cd.add_series("Cumulative students not seated", [round(v) for v in comp["cumulative_unserved"]])
gf = s.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, Inches(0.6), Inches(1.4), Inches(7.4), Inches(5.6), cd)
ch = gf.chart
style_chart(ch, legend=False)
ch.category_axis.reverse_order = True
ch.plots[0].gap_width = 50
ch.plots[0].has_data_labels = True
ch.plots[0].data_labels.number_format, ch.plots[0].data_labels.number_format_is_linked = "#,##0", False
ch.plots[0].data_labels.position = XL_LABEL_POSITION.OUTSIDE_END
ch.plots[0].data_labels.font.size = Pt(12)
ser = ch.series[0]
for i, c in enumerate([GREY, BLUE, BLUE, AQUA, AQUA, ORANGE]):
    pt = ser.points[i]
    pt.format.fill.solid(); pt.format.fill.fore_color.rgb = c
ch.has_title = True
ch.chart_title.text_frame.text = "Cumulative unserved demand over 5 years"
ch.chart_title.text_frame.paragraphs[0].font.size = Pt(14)
r = comp.set_index("scenario")
text(s, 8.4, 1.5, 4.4, 5.5, [
    [("Key trade-offs", {"bold": True, "color": NAVY, "size": 18})],
    f"Status quo leaves {n(r.iloc[0]['cumulative_unserved'])} students unseated over five years.",
    f"50 extra sections a year (Model 2) cut this to {n(r.iloc[1]['cumulative_unserved'])}; 100 a year or Model 1 clear it.",
    f"Raising pass rates by 5 pp shrinks the backlog from {n(r.iloc[1]['final_backlog'])} to "
    f"{n(r.iloc[4]['final_backlog'])}, so academic support also frees capacity.",
    f"With +5%/yr intake growth a fixed budget fails: {n(r.iloc[5]['cumulative_unserved'])} unseated.",
], size=15, color=INK2, space=12)

# 11. Interactive UI ------------------------------------------------------------
s = slide("Interactive UI (Streamlit)")
steps = [("Upload dataset", "Zenodo ZIP, yearly CSVs or processed parquet; counts checked"),
         ("View statistics", "outcomes, demand distribution, top courses, programmes"),
         ("Student flow (SD)", "stocks, flows, transition matrix, what-if sliders"),
         ("Select model", "Model 1 minimum sections or Model 2 limited budget"),
         ("View results", "bottleneck tables, allocations, scenario and sensitivity charts"),
         ("Download", "Excel workbook, CSV per table, cleaned dataset")]
for i, (a, d) in enumerate(steps):
    x = 0.6 + (i % 3) * 4.1
    y = 1.5 + (i // 3) * 2.45
    box(s, x, y, 3.8, 2.1, fill=PANEL, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    box(s, x + 0.25, y + 0.25, 0.6, 0.6, fill=NAVY, shape=MSO_SHAPE.OVAL, label=str(i + 1), size=16, color=WHITE, bold=True)
    text(s, x + 1.0, y + 0.27, 2.7, 0.6, a, size=18, color=NAVY, bold_first=True, anchor=MSO_ANCHOR.MIDDLE)
    text(s, x + 0.25, y + 1.0, 3.4, 1.0, d, size=14, color=INK)
text(s, 0.6, 6.55, 12.1, 0.5, "Sidebar parameters (capacity, growth, threshold, horizon) re-run every model live.   "
     "Run: streamlit run app.py", size=14, color=INK2)

# 12. Conclusions ------------------------------------------------------------
s = slide("Conclusions & Recommendations")
recs = [
    ("Target the bottlenecks", f"{b['bottleneck_courses']} courses in {b['programmes_affected']} programmes need extra sections; "
                               f"{b['extra_sections_to_clear']} sections clear all projected shortfall."),
    ("Rebalance before expanding", f"Model 1 adds {m1['sections_added']} sections and releases {m1['sections_released']}, "
                                   f"a net +{m1['sections_added'] - m1['sections_released']}."),
    ("Spend a limited budget optimally", f"{m2['extra_sections_available']} extra sections remove "
                                         f"{m2['unmet_reduction_pct']:.0f}% of unmet demand; returns fall after about 25."),
    ("Reduce repeat demand", "Higher pass rates shrink the backlog and future demand: repetition is a capacity lever."),
    ("Calibrate capacity", "Bottleneck counts range from 133 to 31 as capacity goes from 30 to 100; use real room and staff data."),
]
for i, (a, d) in enumerate(recs):
    y = 1.45 + i * 1.1
    box(s, 0.6, y, 0.7, 0.7, fill=NAVY, shape=MSO_SHAPE.OVAL, label=str(i + 1), size=18, color=WHITE, bold=True)
    text(s, 1.55, y - 0.05, 11.2, 1.0, [[(a, {"bold": True, "color": NAVY, "size": 18})], d], size=15, color=INK, space=2)

# 13. Team roles ------------------------------------------------------------
s = slide("Team Roles & Responsibilities")
roles = [("Member 1", "Problem formulation, literature review, System Dynamics student-flow model"),
         ("Member 2", "Dataset preprocessing, exploratory analysis, demand and bottleneck identification"),
         ("Member 3", "IP model formulation, solver implementation, scenario and sensitivity analysis"),
         ("Member 4", "Streamlit UI, visualisation, end-to-end integration and testing")]
for i, (a, d) in enumerate(roles):
    x = 0.6 + (i % 2) * 6.15
    y = 1.5 + (i // 2) * 2.75
    box(s, x, y, 5.95, 2.45, fill=PANEL, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    text(s, x + 0.35, y + 0.3, 5.3, 0.6, a, size=22, color=NAVY, bold_first=True, font=TITLE_FONT)
    text(s, x + 0.35, y + 1.0, 5.3, 1.3, d, size=18, color=INK)

out = ROOT / "presentation" / "supplementary_slides.pptx"
prs.save(out)
print("saved", out, len(prs.slides), "slides")
