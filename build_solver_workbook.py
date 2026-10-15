"""Build the full Excel Solver workbook for both optimisation models.

    python build_solver_workbook.py      (run after run_pipeline.py)

Writes outputs/reports/solver_models_full.xlsx with five sheets:

  README              what each sheet is, and Excel Solver's size limits
  Model 1 - all       every course (3,787 rows), optimum by formula, no Solver needed
  Model 1 - Solver    the bottleneck courses as a Solver-ready integer programme
  Model 2 - Solver    all 67 candidate extra sections as a Solver-ready binary programme
  Expected answers    what Solver must return, so the sheets can be checked

Why Model 1 is split across two sheets: Excel's standard Solver accepts at most 200
decision variables, and Model 1 has one per course. Because Model 1 separates by course,
its optimum is the closed form ROUNDUP(demand / capacity), so the full 3,787-course answer
is obtained by formula on the "all" sheet; the Solver sheet demonstrates the same
programme at a size Solver can actually run.
"""
from __future__ import annotations

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from curriculum_planner import config
from curriculum_planner.capacity import modelled_current_sections

CAP = config.DEFAULT_SECTION_CAPACITY
BUDGET = config.DEFAULT_EXTRA_SECTIONS
MAX_PER_PROG = 0          # 0 = no fairness cap; set to e.g. 5 to switch it on
SOLVER_VAR_LIMIT = 200

BOLD = Font(bold=True)
TITLE = Font(bold=True, size=12)
HEAD = Font(bold=True, color="FFFFFF")
HEADFILL = PatternFill("solid", fgColor="1F3864")
CHANGEFILL = PatternFill("solid", fgColor="FFF2CC")     # changing cells, highlighted
TARGETFILL = PatternFill("solid", fgColor="D9EAD3")     # objective cell
THIN = Side(style="thin", color="BFBFBF")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def data():
    bt = pd.read_csv(config.TABLE_DIR / "capacity_bottleneck_table.csv")
    s0 = modelled_current_sections(bt["enrolment"], CAP)
    bt = bt.assign(current_sections_model=s0.astype(int))
    bt["unmet"] = np.maximum(0.0, bt["demand"] - CAP * bt["current_sections_model"])
    return bt


def header(ws, row, labels, widths=None):
    for j, text in enumerate(labels, start=1):
        c = ws.cell(row=row, column=j, value=text)
        c.font = HEAD
        c.fill = HEADFILL
        c.border = BOX
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
    ws.row_dimensions[row].height = 42
    if widths:
        for j, w in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = ws.cell(row=row + 1, column=1)


def steps(ws, row, lines, title="HOW TO RUN IT  (Data > Solver)"):
    ws.cell(row=row, column=1, value=title).font = BOLD
    for i, line in enumerate(lines, start=1):
        ws.cell(row=row + i, column=1, value=line)


# ----------------------------------------------------------------- README
def sheet_readme(wb, bt):
    ws = wb.create_sheet("README")
    ws.column_dimensions["A"].width = 118
    rows = [
        ("University Curriculum Capacity Planning - Excel Solver workbook", TITLE),
        ("", None),
        (f"Section capacity used throughout: {CAP} students per section.", None),
        ("Course identifiers are hashes in the published dataset. They are labels only -", None),
        ("the models use the demand figures, so the hashes do not affect any calculation.", None),
        ("", None),
        ("SHEETS", BOLD),
        ("  Model 1 - all      Every course (3,787 rows). Model 1 separates by course, so the", None),
        ("                     optimum is ROUNDUP(demand / capacity) and is computed by formula.", None),
        ("                     Total = the full answer. No Solver run is needed here.", None),
        ("  Model 1 - Solver   The bottleneck courses set up as an integer programme that Solver", None),
        ("                     can run: changing cells, constraints and objective.", None),
        ("  Model 2 - Solver   All candidate extra sections as a binary (0-1) programme. This is", None),
        ("                     the complete Model 2 - nothing is left out.", None),
        ("  Expected answers   What Solver must return on each sheet.", None),
        ("", None),
        ("WHY MODEL 1 IS SPLIT ACROSS TWO SHEETS", BOLD),
        (f"  Excel's standard Solver accepts at most {SOLVER_VAR_LIMIT} decision variables.", None),
        (f"  Model 1 has one variable per course: {len(bt):,} variables. That is far beyond the limit,", None),
        ("  so Solver cannot optimise every course at once.", None),
        ("  This costs nothing, because each course's constraint involves only its own variable.", None),
        ("  The problem separates, so the optimum per course is exactly ROUNDUP(demand / capacity)", None),
        ("  and the 'all courses' sheet gives the full, exact answer by formula.", None),
        ("", None),
        ("COLOUR KEY", BOLD),
        ("  Yellow cells  = changing cells (the decision variables Solver may alter)", None),
        ("  Green cell    = the objective cell (what Solver minimises or maximises)", None),
        ("", None),
        ("IF SOLVER IS NOT ON THE DATA RIBBON", BOLD),
        ("  File > Options > Add-ins > Manage: Excel Add-ins > Go > tick 'Solver Add-in' > OK.", None),
    ]
    for i, (text, font) in enumerate(rows, start=1):
        c = ws.cell(row=i, column=1, value=text)
        if font:
            c.font = font
    return ws


# ----------------------------------------------------------------- Model 1, all courses
def sheet_m1_all(wb, bt):
    ws = wb.create_sheet("Model 1 - all")
    ws["A1"] = "MODEL 1 - MINIMUM SECTION ALLOCATION, ALL COURSES (closed form, no Solver needed)"
    ws["A1"].font = TITLE
    ws["A2"] = ("Objective: minimise the total number of sections.   "
                "Target: every student gets a seat (cap x sections >= demand).")
    ws["A4"] = "Section capacity (cap)"
    ws["A4"].font = BOLD
    ws["B4"] = CAP
    ws["A5"] = "Demand safety buffer (beta)"
    ws["A5"].font = BOLD
    ws["B5"] = 0.0

    head = 7
    header(ws, head,
           ["Course ID", "Programme", "Demand D_c", "Required seats = (1+beta) x D_c",
            "Sections x_c = ROUNDUP(required / cap)", "Seats provided = cap x x_c",
            "Slack = provided - required", "Current sections", "Change = x_c - current"],
           [16, 16, 12, 15, 16, 14, 14, 11, 12])

    first = head + 1
    for i, r in enumerate(bt.itertuples(index=False)):
        row = first + i
        ws.cell(row=row, column=1, value=r.course_id)
        ws.cell(row=row, column=2, value=r.programme_id)
        ws.cell(row=row, column=3, value=round(float(r.demand), 2))
        ws.cell(row=row, column=4, value=f"=(1+$B$5)*C{row}")
        # a course with no demand runs no section; otherwise at least one, rounded up
        ws.cell(row=row, column=5, value=f"=IF(D{row}<=0,0,MAX(1,ROUNDUP(D{row}/$B$4,0)))")
        ws.cell(row=row, column=6, value=f"=$B$4*E{row}")
        ws.cell(row=row, column=7, value=f"=F{row}-D{row}")
        ws.cell(row=row, column=8, value=int(r.current_sections_model))
        ws.cell(row=row, column=9, value=f"=E{row}-H{row}")
    last = first + len(bt) - 1

    tot = last + 2
    ws.cell(row=tot, column=1, value="TOTAL").font = BOLD
    for col, formula in ((3, f"=SUM(C{first}:C{last})"), (5, f"=SUM(E{first}:E{last})"),
                         (6, f"=SUM(F{first}:F{last})"), (8, f"=SUM(H{first}:H{last})"),
                         (9, f"=SUM(I{first}:I{last})")):
        c = ws.cell(row=tot, column=col, value=formula)
        c.font = BOLD
        c.border = BOX
    ws.cell(row=tot + 1, column=1, value="Courses needing more sections").font = BOLD
    ws.cell(row=tot + 1, column=3, value=f'=COUNTIF(I{first}:I{last},">0")')
    ws.cell(row=tot + 2, column=1, value="Courses that can release a section").font = BOLD
    ws.cell(row=tot + 2, column=3, value=f'=COUNTIF(I{first}:I{last},"<0")')
    ws.cell(row=tot + 3, column=1, value="Mean seat utilisation").font = BOLD
    ws.cell(row=tot + 3, column=3, value=f"=C{tot}/F{tot}")
    ws.cell(row=tot + 3, column=3).number_format = "0.0%"

    steps(ws, tot + 5, [
        "This sheet needs no Solver run.",
        "Each course's constraint involves only its own x_c, so the problem separates by course",
        "and the optimum is the closed form already in column E.",
        "Change the capacity in B4 (or the buffer in B5) and every row and total updates.",
        "To see Solver reproduce this programme at a runnable size, use the 'Model 1 - Solver' sheet.",
    ], title="NOTE")
    return ws


# ----------------------------------------------------------------- Model 1, Solver subset
def sheet_m1_solver(wb, bt):
    sub = bt[bt["unmet"] > 1e-9].sort_values("unmet", ascending=False).reset_index(drop=True)
    ws = wb.create_sheet("Model 1 - Solver")
    ws["A1"] = "MODEL 1 - MINIMUM SECTION ALLOCATION (INTEGER PROGRAMMING), SOLVER-READY"
    ws["A1"].font = TITLE
    ws["A2"] = (f"The {len(sub)} courses that are short at the current allocation. "
                f"Solver's limit is {SOLVER_VAR_LIMIT} variables, so the full {len(bt):,}-course "
                "problem is on the 'Model 1 - all' sheet instead.")
    ws["A4"] = "Section capacity (cap)"
    ws["A4"].font = BOLD
    ws["B4"] = CAP

    head = 6
    header(ws, head,
           ["Course ID", "Programme", "Demand D_c", "Sections x_c  <-- CHANGING CELLS",
            "Seats provided = cap x x_c", "Slack = provided - demand  (must be >= 0)",
            "Current sections", "Change", "Hand check = ROUNDUP(D/cap)"],
           [16, 16, 12, 17, 15, 17, 11, 10, 15])

    first = head + 1
    for i, r in enumerate(sub.itertuples(index=False)):
        row = first + i
        ws.cell(row=row, column=1, value=r.course_id)
        ws.cell(row=row, column=2, value=r.programme_id)
        ws.cell(row=row, column=3, value=round(float(r.demand), 2))
        c = ws.cell(row=row, column=4, value=1)          # changing cell
        c.fill = CHANGEFILL
        c.border = BOX
        ws.cell(row=row, column=5, value=f"=$B$4*D{row}")
        ws.cell(row=row, column=6, value=f"=E{row}-C{row}")
        ws.cell(row=row, column=7, value=int(r.current_sections_model))
        ws.cell(row=row, column=8, value=f"=D{row}-G{row}")
        ws.cell(row=row, column=9, value=f"=MAX(1,ROUNDUP(C{row}/$B$4,0))")
    last = first + len(sub) - 1

    tot = last + 2
    ws.cell(row=tot, column=1, value="TOTAL").font = BOLD
    obj = ws.cell(row=tot, column=4, value=f"=SUM(D{first}:D{last})")
    obj.font = BOLD
    obj.fill = TARGETFILL
    obj.border = BOX
    ws.cell(row=tot, column=7, value=f"=SUM(G{first}:G{last})").font = BOLD
    ws.cell(row=tot, column=9, value=f"=SUM(I{first}:I{last})").font = BOLD
    ws.cell(row=tot + 1, column=1, value="OBJECTIVE CELL (minimise)").font = BOLD
    ws.cell(row=tot + 1, column=4, value=f"=D{tot}")

    steps(ws, tot + 3, [
        f"1. Set Objective:            $D${tot}",
        "2. To:                       Min",
        f"3. By Changing Cells:        $D${first}:$D${last}",
        "4. Subject to the Constraints:",
        f"       $D${first}:$D${last}   int     (sections are whole numbers)",
        f"       $D${first}:$D${last}   >= 1    (a course with demand must run)",
        f"       $F${first}:$F${last}   >= 0    (seats provided >= demand)",
        "5. Select a Solving Method:  Simplex LP",
        "6. Solve.",
        "",
        "CHECK: column D must end up equal to column I, the hand calculation.",
    ])
    return ws


# ----------------------------------------------------------------- Model 2, Solver (complete)
def sheet_m2_solver(wb, bt):
    sub = bt[bt["unmet"] > 1e-9].sort_values("unmet", ascending=False).reset_index(drop=True)
    ws = wb.create_sheet("Model 2 - Solver")
    ws["A1"] = "MODEL 2 - ALLOCATION OF EXTRA SECTIONS (BINARY 0-1 PROGRAMMING), SOLVER-READY"
    ws["A1"].font = TITLE
    ws["A2"] = (f"All {len(sub)} candidate extra sections - this is the complete Model 2, "
                "nothing is left out.")

    ws["A4"] = "Section capacity (cap)"
    ws["A4"].font = BOLD
    ws["B4"] = CAP
    ws["A5"] = "Budget B (extra sections available)"
    ws["A5"].font = BOLD
    ws["B5"] = BUDGET
    ws["A6"] = "Max extra sections per programme (M)"
    ws["A6"].font = BOLD
    ws["B6"] = MAX_PER_PROG if MAX_PER_PROG else 99
    ws["C6"] = "99 = no fairness cap; set to 5 to switch it on"

    head = 8
    header(ws, head,
           ["Course ID", "Programme", "Unmet demand U_c",
            "Students covered if opened  a_c = MIN(cap, U_c)",
            "Open it?  y_c (0 or 1)  <-- CHANGING CELLS", "Students covered = a_c x y_c"],
           [16, 16, 13, 18, 18, 15])

    first = head + 1
    for i, r in enumerate(sub.itertuples(index=False)):
        row = first + i
        ws.cell(row=row, column=1, value=r.course_id)
        ws.cell(row=row, column=2, value=r.programme_id)
        ws.cell(row=row, column=3, value=round(float(r.unmet), 2))
        ws.cell(row=row, column=4, value=f"=MIN($B$4,C{row})")
        c = ws.cell(row=row, column=5, value=0)          # changing cell
        c.fill = CHANGEFILL
        c.border = BOX
        ws.cell(row=row, column=6, value=f"=D{row}*E{row}")
    last = first + len(sub) - 1

    tot = last + 2
    ws.cell(row=tot, column=1, value="TOTAL").font = BOLD
    ws.cell(row=tot, column=3, value=f"=SUM(C{first}:C{last})").font = BOLD
    used = ws.cell(row=tot, column=5, value=f"=SUM(E{first}:E{last})")
    used.font = BOLD
    used.border = BOX
    obj = ws.cell(row=tot, column=6, value=f"=SUMPRODUCT(D{first}:D{last},E{first}:E{last})")
    obj.font = BOLD
    obj.fill = TARGETFILL
    obj.border = BOX

    pr = tot + 2
    ws.cell(row=pr, column=1, value="Programme").font = BOLD
    ws.cell(row=pr, column=2, value="Sections opened").font = BOLD
    ws.cell(row=pr, column=3, value="Limit M").font = BOLD
    progs = sorted(sub["programme_id"].unique())
    for i, pg in enumerate(progs, start=1):
        row = pr + i
        ws.cell(row=row, column=1, value=pg)
        ws.cell(row=row, column=2,
                value=f"=SUMIF($B${first}:$B${last},A{row},$E${first}:$E${last})").border = BOX
        ws.cell(row=row, column=3, value="=$B$6")
    pf, pl = pr + 1, pr + len(progs)

    res = pl + 2
    ws.cell(row=res, column=1, value="OBJECTIVE CELL (maximise)").font = BOLD
    ws.cell(row=res, column=2, value=f"=F{tot}")
    ws.cell(row=res + 1, column=1, value="Sections used").font = BOLD
    ws.cell(row=res + 1, column=2, value=f"=E{tot}")
    ws.cell(row=res + 2, column=1, value="Unmet demand remaining").font = BOLD
    ws.cell(row=res + 2, column=2, value=f"=C{tot}-F{tot}")

    steps(ws, res + 4, [
        f"1. Set Objective:            $F${tot}",
        "2. To:                       Max",
        f"3. By Changing Cells:        $E${first}:$E${last}",
        "4. Subject to the Constraints:",
        f"       $E${first}:$E${last}   bin                (open it or not)",
        f"       $E${tot}  <= $B$5                         (budget: at most B sections)",
        f"       $B${pf}:$B${pl}  <= $B$6                  (at most M per programme)",
        "5. Select a Solving Method:  Simplex LP",
        "6. Solve.",
        "",
        "Set B6 to 5 and re-solve to see what the fairness cap costs.",
    ])
    return ws


# ----------------------------------------------------------------- expected answers
def sheet_answers(wb, bt):
    from curriculum_planner.optimization import solve_extra_sections, solve_min_sections

    sub = bt[bt["unmet"] > 1e-9]
    s0 = bt["current_sections_model"].to_numpy()
    r1 = solve_min_sections(bt[["course_id", "programme_id", "demand"]], CAP)
    change = r1.allocation["sections"].to_numpy() - s0
    n_more, n_less = int((change > 0).sum()), int((change < 0).sum())
    r2 = solve_extra_sections(bt, s0, BUDGET, CAP)
    r2f = solve_extra_sections(bt, s0, BUDGET, CAP, max_per_programme=5)
    x_sub = np.maximum(1, np.ceil(sub["demand"] / CAP)).astype(int)

    ws = wb.create_sheet("Expected answers")
    ws.column_dimensions["A"].width = 52
    ws.column_dimensions["B"].width = 16
    ws.column_dimensions["C"].width = 56
    ws["A1"] = "EXPECTED ANSWERS - Solver must reproduce these"
    ws["A1"].font = TITLE

    rows = [
        ("", "", ""),
        ("MODEL 1 - all courses sheet", "", ""),
        ("Total sections (optimum)", int(r1.summary["total_sections"]), "closed form, no Solver run"),
        ("Modelled current sections", int(s0.sum()), ""),
        ("Net change", int(r1.summary["total_sections"] - s0.sum()), ""),
        ("Courses needing more sections", n_more, ""),
        ("Courses that can release", n_less, ""),
        ("Mean seat utilisation", round(float(r1.summary["mean_utilisation"]), 3), ""),
        ("", "", ""),
        ("MODEL 1 - Solver sheet", "", ""),
        (f"Total sections over the {len(sub)} short courses", int(x_sub.sum()),
         "Solver must match the hand-check column"),
        ("Their current sections", int(sub["current_sections_model"].sum()), ""),
        ("", "", ""),
        (f"MODEL 2 - Solver sheet (B = {BUDGET})", "", ""),
        ("Unmet demand before", round(float(r2.summary["unmet_demand_before"]), 2), "students"),
        ("Students covered (objective)", round(float(r2.summary["demand_covered"]), 2), "maximised"),
        ("Unmet demand after", round(float(r2.summary["unmet_demand_after"]), 2), "students"),
        ("Sections used", int(r2.summary["extra_sections_used"]), f"of {BUDGET} available"),
        ("Courses still short", int(r2.summary["courses_short_after"]),
         f"was {int(r2.summary['courses_short_before'])}"),
        ("", "", ""),
        ("MODEL 2 with fairness cap M = 5", "", ""),
        ("Unmet demand after", round(float(r2f.summary["unmet_demand_after"]), 2),
         "the cap costs "
         f"{r2f.summary['unmet_demand_after'] - r2.summary['unmet_demand_after']:.2f} students"),
    ]
    for i, (a, b, c) in enumerate(rows, start=2):
        ws.cell(row=i, column=1, value=a).font = BOLD if a and b == "" else Font()
        ws.cell(row=i, column=2, value=b)
        ws.cell(row=i, column=3, value=c)
    return ws


def main():
    bt = data()
    wb = Workbook()
    wb.remove(wb.active)
    sheet_readme(wb, bt)
    sheet_m1_all(wb, bt)
    sheet_m1_solver(wb, bt)
    sheet_m2_solver(wb, bt)
    sheet_answers(wb, bt)
    out = config.REPORT_DIR / "solver_models_full.xlsx"
    wb.save(out)
    print(f"wrote {out}")
    print(f"  Model 1 - all      : {len(bt):,} course rows")
    print(f"  Model 1 - Solver   : {int((bt['unmet'] > 1e-9).sum())} variables")
    print(f"  Model 2 - Solver   : {int((bt['unmet'] > 1e-9).sum())} binary variables, "
          f"{bt.loc[bt['unmet'] > 1e-9, 'programme_id'].nunique()} programme constraints")


if __name__ == "__main__":
    main()
